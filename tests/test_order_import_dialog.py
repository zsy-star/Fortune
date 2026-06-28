from __future__ import annotations

import os
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QMessageBox
from sqlalchemy import func, select

from models import Order
from services.log_service import LogService
from services.order_import_service import OrderImportService
from services.order_intake_service import OrderIntakeService
from services.order_service import OrderService
from services.settings_service import SettingsService
from ui.app_events import app_events
from ui.dialogs.order_import_dialog import OrderImportDialog
from ui.pages.order_detail_page import OrderDetailPage


def app() -> QApplication:
    return QApplication.instance() or QApplication([])


def count_orders(session_factory) -> int:
    with session_factory() as session:
        return int(session.scalar(select(func.count(Order.id))) or 0)


def count_import_logs(session_factory) -> int:
    return LogService(session_factory).count_logs(action="order/import")


def save_order_text(
    session_factory,
    text: str,
    *,
    region: str = "澳门",
    customer_name: str | None = None,
    channel: str = "微信",
):
    intake = OrderIntakeService(session_factory)
    preview = intake.preview_raw_text(text, region=region, customer_name=customer_name, channel=channel)
    result = intake.save_preview(preview)
    assert result.success
    assert result.order is not None
    return result.order


class EmptySettingsService:
    def list_declarers(self):
        return []


class BrokenSettingsService:
    def list_declarers(self):
        raise RuntimeError("settings failed")


def test_order_import_service_txt_reads_non_empty_lines() -> None:
    service = OrderImportService()

    result = service.extract_txt_lines("\n01各10\n  \n02各5\n")

    assert result.lines == ["01各10", "02各5"]
    assert result.skipped_count == 2


def test_order_import_service_csv_with_header_reads_text_column() -> None:
    service = OrderImportService()

    result = service.extract_csv_lines("name,订单内容\n甲,01各10\n乙,无效文本\n")

    assert result.lines == ["01各10", "无效文本"]


def test_order_import_service_csv_without_header_uses_first_column() -> None:
    service = OrderImportService()

    result = service.extract_csv_lines("01各10,备注\n02各5,备注\n")

    assert result.lines == ["01各10", "02各5"]


def test_order_import_service_creates_txt_csv_and_xlsx_templates(tmp_path) -> None:
    service = OrderImportService()
    txt_path = tmp_path / "template.txt"
    csv_path = tmp_path / "template.csv"
    xlsx_path = tmp_path / "template.xlsx"

    service.create_template(txt_path)
    service.create_template(csv_path)
    service.create_template(xlsx_path)

    assert txt_path.read_text(encoding="utf-8") == "01各10\n02各5\n"
    csv_lines = csv_path.read_text(encoding="utf-8").splitlines()
    assert csv_lines[0] == "订单文本"
    assert "01各10" in csv_lines[1]

    workbook = load_workbook(xlsx_path, data_only=True)
    assert workbook.active.title == "订单导入模板"
    assert workbook.active["A1"].value == "订单文本"
    workbook.close()
    result = service.extract_xlsx_lines(xlsx_path)
    assert result.lines == ["01各10", "02各5"]


def write_xlsx(path, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_order_import_service_xlsx_with_header_reads_order_text_column(tmp_path) -> None:
    service = OrderImportService()
    path = tmp_path / "orders.xlsx"
    write_xlsx(path, [["备注", "订单文本"], ["A", "01各10"], ["B", "02各5"]])

    result = service.read_file(path)

    assert result.error == ""
    assert result.lines == ["01各10", "02各5"]


def test_order_import_service_xlsx_without_header_uses_first_column(tmp_path) -> None:
    service = OrderImportService()
    path = tmp_path / "orders.xlsx"
    write_xlsx(path, [["01各10", "备注"], ["02各5", "备注"]])

    result = service.read_file(path)

    assert result.lines == ["01各10", "02各5"]


def test_order_import_service_xlsx_skips_empty_rows(tmp_path) -> None:
    service = OrderImportService()
    path = tmp_path / "orders.xlsx"
    write_xlsx(path, [["订单内容"], [None], ["01各10"], ["", ""], ["02各5"]])

    result = service.read_file(path)

    assert result.lines == ["01各10", "02各5"]
    assert result.skipped_count == 2


def test_order_import_service_xlsx_read_failure_returns_clear_error(tmp_path) -> None:
    service = OrderImportService()
    path = tmp_path / "broken.xlsx"
    path.write_text("not a real xlsx", encoding="utf-8")

    result = service.read_file(path)

    assert result.lines == []
    assert "xlsx 文件读取失败" in result.error


def test_order_import_service_rejects_unsupported_file_type(tmp_path) -> None:
    service = OrderImportService()
    path = tmp_path / "orders.docx"
    path.write_text("01各10", encoding="utf-8")

    result = service.read_file(path)

    assert result.lines == []
    assert "当前仅支持 txt/csv/xlsx" in result.error


def test_order_import_service_preview_success_and_failure_rows() -> None:
    service = OrderImportService()

    preview = service.preview_lines(["01各10", "明显无效"], region_mode="澳门")

    assert preview.total_count == 2
    assert preview.success_count == 1
    assert preview.failure_count == 1
    assert preview.rows[0].success is True
    assert preview.rows[0].region == "澳门"
    assert preview.rows[0].bet_type_summary == "纯数字"
    assert preview.rows[0].amount_total == 10
    assert preview.rows[0].item_count == 1
    assert preview.rows[1].success is False
    assert preview.rows[1].error


def test_order_import_service_marks_history_duplicate_read_only(session_factory) -> None:
    existing = save_order_text(session_factory, "01各10", region="澳门", customer_name="林林")
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    before = count_orders(session_factory)

    preview = service.preview_lines(
        ["０１各10", "02各5"],
        region_mode="澳门",
        history_duplicate_check=True,
        customer_name="林林",
    )

    assert count_orders(session_factory) == before
    assert preview.history_duplicate_check_enabled is True
    assert preview.history_duplicate_count == 1
    assert preview.rows[0].history_duplicate_warning
    assert preview.rows[0].history_duplicate_order_id == existing.id
    assert preview.rows[0].history_duplicate_order_no == existing.order_no
    assert not preview.rows[1].history_duplicate_warning


def test_order_import_history_duplicate_detection_degrades_safely() -> None:
    class BrokenIntake:
        _session_factory = None

    class FakeLogService:
        def create_log(self, **kwargs):
            return SimpleNamespace(id=1)

    service = OrderImportService(order_intake_service=BrokenIntake(), log_service=FakeLogService())

    preview = service.preview_lines(["01各10"], region_mode="澳门", history_duplicate_check=True)

    assert preview.history_duplicate_check_enabled is False
    assert "历史重复检测失败" in preview.history_duplicate_error


def test_order_import_dialog_initializes_with_confirm_button_disabled() -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=EmptySettingsService())

    assert dialog.windowTitle() == "订单导入预览"
    assert dialog._btn_select_file.text() == "选择文件"
    assert dialog._btn_download_template.text() == "下载导入模板"
    assert dialog._btn_preview.text() == "重新解析"
    assert dialog._btn_confirm_import.text() == "确认导入成功行"
    assert not dialog._btn_confirm_import.isEnabled()
    assert dialog._btn_copy_errors.text() == "复制错误报告"
    assert dialog._btn_copy_result.text() == "复制导入结果"
    assert dialog._btn_export_result.text() == "导出导入结果"
    assert dialog._btn_close.text() == "关闭"
    assert dialog._declarer_combo.currentText() == "未设置申报人"
    assert dialog._channel_combo.currentText() == "导入"
    assert dialog._channel_combo.findText("XLSX导入") >= 0
    assert dialog._chk_history_duplicate.text() == "检查历史重复订单"
    assert dialog._chk_history_duplicate.isChecked()
    assert dialog._chk_skip_history_duplicate.text() == "导入时跳过历史疑似重复行"
    assert dialog._chk_skip_history_duplicate.isChecked()
    assert "尚未确认导入" in dialog._stats_label.text()


def test_order_import_dialog_file_picker_supports_xlsx(monkeypatch) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=EmptySettingsService())
    captured = {}

    def fake_get_open_file_name(*args, **kwargs):
        captured["filter"] = args[3]
        return "", ""

    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getOpenFileName",
        fake_get_open_file_name,
    )

    dialog._on_select_file()

    assert "*.xlsx" in captured["filter"]


def test_order_import_dialog_template_button_cancel_success_and_failure(
    session_factory, tmp_path, monkeypatch
) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=EmptySettingsService())
    before = count_orders(session_factory)

    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: ("", ""),
    )
    dialog._btn_download_template.click()
    assert "已取消生成导入模板" in dialog._stats_label.text()

    txt_path = tmp_path / "manual_template.txt"
    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(txt_path), "Text Files (*.txt)"),
    )
    with patch("ui.dialogs.order_import_dialog.QMessageBox.information") as info:
        dialog._btn_download_template.click()
    assert info.called
    assert txt_path.read_text(encoding="utf-8") == "01各10\n02各5\n"
    assert "已生成导入模板" in dialog._stats_label.text()
    assert count_orders(session_factory) == before

    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(tmp_path / "missing" / "template.txt"), "Text Files (*.txt)"),
    )
    with patch("ui.dialogs.order_import_dialog.QMessageBox.warning") as warning:
        dialog._btn_download_template.click()
    assert warning.called
    assert "生成导入模板失败" in dialog._stats_label.text()


def test_order_import_dialog_generates_csv_and_xlsx_templates(tmp_path, monkeypatch) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=EmptySettingsService())
    csv_path = tmp_path / "template.csv"
    xlsx_path = tmp_path / "template.xlsx"

    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(csv_path), "CSV Files (*.csv)"),
    )
    with patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._on_download_template()
    assert csv_path.read_text(encoding="utf-8").splitlines()[0] == "订单文本"

    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(xlsx_path), "Excel Files (*.xlsx)"),
    )
    with patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._on_download_template()
    assert dialog._import_service.extract_xlsx_lines(xlsx_path).lines == ["01各10", "02各5"]


def test_order_import_dialog_loads_declarers_and_plan_from_settings(session_factory) -> None:
    app()
    settings = SettingsService(session_factory)
    plan = settings.create_plan("47倍4水")
    settings.add_declarer("林林", plan.id)

    dialog = OrderImportDialog(
        import_service=OrderImportService(order_intake_service=OrderIntakeService(session_factory)),
        settings_service=settings,
    )
    index = dialog._declarer_combo.findText("林林")

    assert index >= 0
    dialog._declarer_combo.setCurrentIndex(index)
    assert dialog._selected_customer_name() == "林林"
    assert dialog._selected_config_plan_name() == "47倍4水"
    assert dialog._plan_label.text() == "配置方案：47倍4水"


def test_order_import_dialog_declarer_settings_failure_falls_back() -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=BrokenSettingsService())

    assert dialog._declarer_combo.currentText() == "未设置申报人"
    assert dialog._selected_customer_name() is None
    assert dialog._plan_label.text() == "配置方案：配置读取失败"
    assert "申报人配置读取失败" in dialog._stats_label.text()


def test_order_import_dialog_loads_txt_and_updates_preview_table(tmp_path) -> None:
    app()
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n\n明显无效\n", encoding="utf-8")
    dialog = OrderImportDialog(import_service=OrderImportService())

    dialog.load_file(str(path))

    assert dialog._path_edit.text() == str(path)
    assert dialog._table.rowCount() == 2
    assert dialog._table.item(0, 3).text() == "成功"
    assert dialog._table.item(1, 3).text() == "失败"
    assert dialog._table.item(1, 7).text()
    assert "总行数：2" in dialog._stats_label.text()
    assert "解析成功：1" in dialog._stats_label.text()
    assert "解析失败：1" in dialog._stats_label.text()
    assert "跳过空行：1" in dialog._stats_label.text()
    assert dialog._btn_confirm_import.isEnabled()


def test_order_import_dialog_loads_csv_header_and_no_header(tmp_path) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService())
    header_path = tmp_path / "orders_header.csv"
    header_path.write_text("备注,order_text\nA,01各10\n", encoding="utf-8")
    no_header_path = tmp_path / "orders_no_header.csv"
    no_header_path.write_text("02各5,备注\n", encoding="utf-8")

    dialog.load_file(str(header_path))
    assert dialog._raw_text.toPlainText() == "01各10"
    assert dialog._table.item(0, 3).text() == "成功"

    dialog.load_file(str(no_header_path))
    assert dialog._raw_text.toPlainText() == "02各5"
    assert dialog._table.item(0, 3).text() == "成功"


def test_order_import_dialog_loads_xlsx_and_suggests_channel(tmp_path) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=EmptySettingsService())
    path = tmp_path / "orders.xlsx"
    write_xlsx(path, [["投注内容"], ["01各10"], ["明显无效"]])

    dialog.load_file(str(path))

    assert dialog._path_edit.text() == str(path)
    assert dialog._channel_combo.currentText() == "XLSX导入"
    assert dialog._raw_text.toPlainText() == "01各10\n明显无效"
    assert dialog._table.rowCount() == 2
    assert dialog._table.item(0, 3).text() == "成功"
    assert dialog._table.item(1, 3).text() == "失败"
    assert dialog._btn_confirm_import.isEnabled()


def test_order_import_dialog_xlsx_does_not_override_manual_channel(tmp_path) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=EmptySettingsService())
    dialog._channel_combo.setCurrentText("手工整理导入")
    path = tmp_path / "orders.xlsx"
    write_xlsx(path, [["订单内容"], ["01各10"]])

    dialog.load_file(str(path))

    assert dialog._channel_combo.currentText() == "手工整理导入"


def test_order_import_dialog_xlsx_read_failure_shows_error(tmp_path) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService(), settings_service=EmptySettingsService())
    path = tmp_path / "broken.xlsx"
    path.write_text("not a real xlsx", encoding="utf-8")

    with patch("ui.dialogs.order_import_dialog.QMessageBox.warning") as warning:
        dialog.load_file(str(path))

    assert warning.called
    assert "xlsx 文件读取失败" in dialog._stats_label.text()
    assert dialog._table.rowCount() == 0


def test_order_import_dialog_unsupported_type_shows_error(tmp_path) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService())
    path = tmp_path / "orders.docx"
    path.write_text("01各10", encoding="utf-8")

    with patch("ui.dialogs.order_import_dialog.QMessageBox.warning") as warning:
        dialog.load_file(str(path))

    assert warning.called
    assert "当前仅支持 txt/csv/xlsx" in dialog._stats_label.text()
    assert dialog._table.rowCount() == 0


def test_order_import_dialog_copy_error_report_handles_errors_and_no_errors(tmp_path) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService())
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n明显无效\n", encoding="utf-8")
    dialog.load_file(str(path))

    dialog._btn_copy_errors.click()
    assert "明显无效" in QApplication.clipboard().text()
    assert "已复制错误报告" in dialog._stats_label.text()

    dialog._raw_text.setPlainText("01各10")
    dialog._on_preview()
    dialog._btn_copy_errors.click()
    assert QApplication.clipboard().text() == "当前没有解析失败行"


def test_order_import_preview_does_not_write_database(session_factory, tmp_path) -> None:
    app()
    before = count_orders(session_factory)
    dialog = OrderImportDialog(import_service=OrderImportService())
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n02各5\n", encoding="utf-8")

    dialog.load_file(str(path))
    dialog._on_preview()
    dialog._btn_copy_errors.click()

    assert count_orders(session_factory) == before


def test_order_import_confirm_cancel_does_not_write_database(session_factory, tmp_path) -> None:
    app()
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service)
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n", encoding="utf-8")
    dialog.load_file(str(path))
    before = count_orders(session_factory)

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.No,
    ) as question:
        dialog._btn_confirm_import.click()

    assert question.called
    assert count_orders(session_factory) == before
    assert "已取消确认导入" in dialog._stats_label.text()


def test_order_import_confirm_saves_only_success_rows_and_writes_log(
    session_factory, tmp_path
) -> None:
    app()
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service)
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n\n明显无效\n02各5\n", encoding="utf-8")
    dialog.load_file(str(path))
    before = count_orders(session_factory)
    before_logs = count_import_logs(session_factory)
    events = {"orders": 0, "logs": 0}

    def on_orders():
        events["orders"] += 1

    def on_logs():
        events["logs"] += 1

    app_events.orders_changed.connect(on_orders)
    app_events.logs_changed.connect(on_logs)
    try:
        with patch(
            "ui.dialogs.order_import_dialog.QMessageBox.question",
            return_value=QMessageBox.StandardButton.Yes,
        ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
            dialog._btn_confirm_import.click()
    finally:
        app_events.orders_changed.disconnect(on_orders)
        app_events.logs_changed.disconnect(on_logs)

    assert count_orders(session_factory) == before + 2
    assert count_import_logs(session_factory) == before_logs + 1
    assert dialog._table.item(0, 8).text() == "已导入"
    assert dialog._table.item(1, 8).text() == "未导入，解析失败"
    assert dialog._table.item(2, 8).text() == "已导入"
    assert "本次应导入：2" in dialog._stats_label.text()
    assert "成功导入：2" in dialog._stats_label.text()
    assert "保存失败：0" in dialog._stats_label.text()
    assert "跳过失败解析行：1" in dialog._stats_label.text()
    assert not dialog._btn_confirm_import.isEnabled()
    assert events == {"orders": 1, "logs": 1}

    logs = LogService(session_factory).list_logs(action="order/import")
    assert logs
    description = logs[0].description
    assert "total_rows=3" in description
    assert "parse_success=2" in description
    assert "parse_failed=1" in description
    assert "imported=2" in description
    assert "save_failed=0" in description
    assert "skipped_empty=1" in description


def test_order_import_confirm_passes_declarer_channel_and_plan_to_save_log_and_report(
    session_factory, tmp_path, monkeypatch
) -> None:
    app()
    settings = SettingsService(session_factory)
    plan = settings.create_plan("46倍6水")
    settings.add_declarer("老汪", plan.id)
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service, settings_service=settings)
    dialog._region_combo.setCurrentText("澳门")
    dialog._declarer_combo.setCurrentIndex(dialog._declarer_combo.findText("老汪"))
    dialog._channel_combo.setCurrentText("手工整理导入")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n", encoding="utf-8")
    dialog.load_file(str(path))

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    with session_factory() as session:
        order = session.scalar(select(Order))
        assert order is not None
        assert order.customer_name == "老汪"
        assert order.channel == "手工整理导入"

    description = LogService(session_factory).list_logs(action="order/import")[0].description
    assert "declarer=老汪" in description
    assert "channel=手工整理导入" in description
    assert "config_plan=46倍6水" in description

    report = dialog._build_import_result_report()
    assert "申报人：老汪" in report
    assert "渠道：手工整理导入" in report
    assert "配置方案：46倍6水" in report

    csv_path = tmp_path / "import_report.csv"
    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(csv_path), "CSV Files (*.csv)"),
    )
    dialog._on_export_import_result()
    lines = csv_path.read_text(encoding="utf-8").splitlines()
    assert "申报人,渠道,配置方案" in lines[0]
    assert "老汪,手工整理导入,46倍6水" in lines[1]


def test_order_import_xlsx_confirm_uses_existing_save_flow_and_report_channel(
    session_factory, tmp_path
) -> None:
    app()
    intake = OrderIntakeService(session_factory)
    service = OrderImportService(order_intake_service=intake)
    dialog = OrderImportDialog(import_service=service, settings_service=EmptySettingsService())
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.xlsx"
    write_xlsx(path, [["订单文本"], ["01各10"], ["明显无效"], ["02各5"]])
    dialog.load_file(str(path))
    before = count_orders(session_factory)

    with patch.object(intake, "save_preview", wraps=intake.save_preview) as save_preview, patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    assert save_preview.called
    assert count_orders(session_factory) == before + 2
    assert dialog._table.item(0, 8).text() == "已导入"
    assert dialog._table.item(1, 8).text() == "未导入，解析失败"
    assert dialog._table.item(2, 8).text() == "已导入"
    report = dialog._build_import_result_report()
    assert "渠道：XLSX导入" in report
    assert "成功导入数：2" in report
    assert "解析失败数：1" in report


def test_order_import_result_report_copy_and_export(session_factory, tmp_path, monkeypatch) -> None:
    app()
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service)
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n明显无效\n", encoding="utf-8")
    dialog.load_file(str(path))

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    report = dialog._build_import_result_report()
    assert "订单导入结果报告" in report
    assert "总行数：2" in report
    assert "解析成功数：1" in report
    assert "解析失败数：1" in report
    assert "成功导入数：1" in report
    assert "保存失败数：0" in report
    assert "失败行未导入" in report
    assert "本次未执行结算" in report
    assert "本次未计算赔付、余额、返水" in report

    dialog._btn_copy_result.click()
    assert "订单导入结果报告" in QApplication.clipboard().text()
    assert "已复制导入结果报告到剪贴板" in dialog._stats_label.text()

    txt_path = tmp_path / "import_report.txt"
    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(txt_path), "Text Files (*.txt)"),
    )
    dialog._on_export_import_result()
    assert txt_path.read_text(encoding="utf-8").startswith("订单导入结果报告")

    csv_path = tmp_path / "import_report.csv"
    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(csv_path), "CSV Files (*.csv)"),
    )
    dialog._on_export_import_result()
    first_line = csv_path.read_text(encoding="utf-8").splitlines()[0]
    assert "行号,原始文本,解析状态,导入状态,地区,申报人,渠道,配置方案,投注类型,金额,条目数,历史重复提示,历史订单,是否跳过,错误原因" in first_line


def test_order_import_result_copy_before_preview_and_export_cancel_or_failure(tmp_path, monkeypatch) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService())

    dialog._btn_copy_result.click()
    assert "当前暂无导入结果" in dialog._stats_label.text()

    path = tmp_path / "orders.txt"
    path.write_text("01各10\n", encoding="utf-8")
    dialog.load_file(str(path))

    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: ("", ""),
    )
    dialog._on_export_import_result()
    assert "已取消导出导入结果" in dialog._stats_label.text()

    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(tmp_path), "Text Files (*.txt)"),
    )
    dialog._on_export_import_result()
    assert "导出导入结果失败" in dialog._stats_label.text()


def test_order_import_marks_duplicate_rows_and_confirmation_warns(tmp_path) -> None:
    app()
    dialog = OrderImportDialog(import_service=OrderImportService())
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n 01各10 \n02各5\n", encoding="utf-8")
    dialog.load_file(str(path))

    assert dialog._preview.rows[0].duplicate_warning == "疑似重复行"
    assert dialog._preview.rows[1].duplicate_warning == "疑似重复行"
    assert "疑似重复行" in dialog._table.item(0, 7).text()

    captured = {}

    def fake_question(_parent, _title, text, *_args):
        captured["text"] = text
        return QMessageBox.StandardButton.No

    with patch("ui.dialogs.order_import_dialog.QMessageBox.question", side_effect=fake_question):
        dialog._btn_confirm_import.click()

    assert "当前预览中存在疑似重复行" in captured["text"]


def test_order_import_dialog_marks_history_duplicate_and_confirmation_warns(
    session_factory, tmp_path
) -> None:
    app()
    save_order_text(session_factory, "01各10", region="澳门")
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service, settings_service=EmptySettingsService())
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n02各5\n", encoding="utf-8")
    dialog.load_file(str(path))

    assert dialog._preview.history_duplicate_count == 1
    assert "疑似历史重复" in dialog._table.item(0, 7).text()
    assert dialog._table.item(0, 8).text() == "待跳过，疑似历史重复"
    assert "历史疑似重复：1" in dialog._stats_label.text()

    captured = {}

    def fake_question(_parent, _title, text, *_args):
        captured["text"] = text
        return QMessageBox.StandardButton.No

    with patch("ui.dialogs.order_import_dialog.QMessageBox.question", side_effect=fake_question):
        dialog._btn_confirm_import.click()

    assert "历史疑似重复行数：1" in captured["text"]
    assert "导入时跳过历史疑似重复行：是" in captured["text"]
    assert "当前预览中存在历史疑似重复行" in captured["text"]


def test_order_import_skips_history_duplicates_when_checked(session_factory, tmp_path) -> None:
    app()
    existing = save_order_text(session_factory, "01各10", region="澳门")
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service, settings_service=EmptySettingsService())
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n明显无效\n02各5\n", encoding="utf-8")
    dialog.load_file(str(path))
    before = count_orders(session_factory)

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    assert count_orders(session_factory) == before + 1
    assert dialog._table.item(0, 8).text() == "已跳过，疑似历史重复"
    assert dialog._table.item(1, 8).text() == "未导入，解析失败"
    assert dialog._table.item(2, 8).text() == "已导入"
    assert "跳过历史重复：1" in dialog._stats_label.text()

    description = LogService(session_factory).list_logs(action="order/import")[0].description
    assert f"history_duplicate_count=1" in description
    assert "skipped_history_duplicate=1" in description
    assert dialog._preview.rows[0].history_duplicate_order_no == existing.order_no


def test_order_import_allows_history_duplicates_when_skip_unchecked(session_factory, tmp_path) -> None:
    app()
    save_order_text(session_factory, "01各10", region="澳门")
    intake = OrderIntakeService(session_factory)
    service = OrderImportService(order_intake_service=intake)
    dialog = OrderImportDialog(import_service=service, settings_service=EmptySettingsService())
    dialog._region_combo.setCurrentText("澳门")
    dialog._chk_skip_history_duplicate.setChecked(False)
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n", encoding="utf-8")
    dialog.load_file(str(path))
    before = count_orders(session_factory)

    with patch.object(intake, "save_preview", wraps=intake.save_preview) as save_preview, patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    assert save_preview.called
    assert count_orders(session_factory) == before + 1
    assert dialog._table.item(0, 8).text() == "已导入"
    assert "跳过历史重复：0" in dialog._stats_label.text()


def test_order_import_report_and_csv_include_history_duplicate_fields(
    session_factory, tmp_path, monkeypatch
) -> None:
    app()
    save_order_text(session_factory, "01各10", region="澳门")
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service, settings_service=EmptySettingsService())
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n02各5\n", encoding="utf-8")
    dialog.load_file(str(path))

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    report = dialog._build_import_result_report()
    assert "历史重复检测：已启用" in report
    assert "历史疑似重复行数：1" in report
    assert "跳过历史重复行数：1" in report
    assert "疑似历史重复" in report

    csv_path = tmp_path / "import_report.csv"
    monkeypatch.setattr(
        "ui.dialogs.order_import_dialog.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(csv_path), "CSV Files (*.csv)"),
    )
    dialog._on_export_import_result()
    first_line = csv_path.read_text(encoding="utf-8").splitlines()[0]
    assert "历史重复提示,历史订单,是否跳过" in first_line


def test_order_import_reselect_file_resets_import_state(session_factory, tmp_path) -> None:
    app()
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service)
    dialog._region_combo.setCurrentText("澳门")
    first = tmp_path / "first.txt"
    first.write_text("01各10\n", encoding="utf-8")
    second = tmp_path / "second.txt"
    second.write_text("02各5\n", encoding="utf-8")
    dialog.load_file(str(first))

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    assert not dialog._btn_confirm_import.isEnabled()
    dialog.load_file(str(second))
    assert dialog._btn_confirm_import.isEnabled()
    assert dialog._table.item(0, 8).text() == "未导入"


def test_order_import_reparse_and_reselect_keep_declarer_and_channel(session_factory, tmp_path) -> None:
    app()
    settings = SettingsService(session_factory)
    plan = settings.create_plan("47倍4水")
    settings.add_declarer("林林", plan.id)
    dialog = OrderImportDialog(
        import_service=OrderImportService(order_intake_service=OrderIntakeService(session_factory)),
        settings_service=settings,
    )
    dialog._declarer_combo.setCurrentIndex(dialog._declarer_combo.findText("林林"))
    dialog._channel_combo.setCurrentText("手工整理导入")
    first = tmp_path / "first.txt"
    first.write_text("01各10\n", encoding="utf-8")
    second = tmp_path / "second.csv"
    second.write_text("order_text\n02各5\n", encoding="utf-8")

    dialog.load_file(str(first))
    dialog._raw_text.setPlainText("03各5")
    dialog._on_preview()
    assert dialog._declarer_combo.currentText() == "林林"
    assert dialog._channel_combo.currentText() == "手工整理导入"

    dialog.load_file(str(second))
    assert dialog._declarer_combo.currentText() == "林林"
    assert dialog._channel_combo.currentText() == "手工整理导入"
    assert dialog._selected_config_plan_name() == "47倍4水"


def test_order_import_confirm_all_success_and_prevents_duplicate_click(
    session_factory, tmp_path
) -> None:
    app()
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service)
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n02各5\n", encoding="utf-8")
    dialog.load_file(str(path))

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()
        after_first = count_orders(session_factory)
        dialog._channel_combo.setCurrentText("CSV导入")
        dialog._btn_confirm_import.click()

    assert count_orders(session_factory) == after_first
    assert "成功导入：2" in dialog._stats_label.text()
    assert not dialog._btn_confirm_import.isEnabled()


def test_order_import_reparse_resets_import_state(session_factory, tmp_path) -> None:
    app()
    service = OrderImportService(order_intake_service=OrderIntakeService(session_factory))
    dialog = OrderImportDialog(import_service=service)
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n", encoding="utf-8")
    dialog.load_file(str(path))

    with patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    assert not dialog._btn_confirm_import.isEnabled()
    dialog._raw_text.setPlainText("02各5")
    dialog._on_preview()
    assert dialog._btn_confirm_import.isEnabled()
    assert dialog._table.item(0, 8).text() == "未导入"


def test_order_import_save_failure_is_shown_without_crashing() -> None:
    app()

    class BrokenIntake:
        _session_factory = None

        def preview_raw_text(self, *args, **kwargs):
            raise RuntimeError("模拟保存失败")

    class FakeLogService:
        def create_log(self, **kwargs):
            return SimpleNamespace(id=1)

    service = OrderImportService(order_intake_service=BrokenIntake(), log_service=FakeLogService())
    preview = service.preview_lines(["01各10"], region_mode="澳门")

    result = service.confirm_import(preview, file_path="orders.txt")

    assert result.imported_count == 0
    assert result.save_failed_count == 1
    assert result.preview.rows[0].import_status == "导入失败"
    assert "模拟保存失败" in result.preview.rows[0].import_error


def test_order_import_confirm_uses_order_intake_service_save_preview(
    session_factory, tmp_path
) -> None:
    app()
    intake = OrderIntakeService(session_factory)
    service = OrderImportService(order_intake_service=intake)
    dialog = OrderImportDialog(import_service=service)
    dialog._region_combo.setCurrentText("澳门")
    path = tmp_path / "orders.txt"
    path.write_text("01各10\n", encoding="utf-8")
    dialog.load_file(str(path))

    with patch.object(intake, "save_preview", wraps=intake.save_preview) as save_preview, patch(
        "ui.dialogs.order_import_dialog.QMessageBox.question",
        return_value=QMessageBox.StandardButton.Yes,
    ), patch("ui.dialogs.order_import_dialog.QMessageBox.information"):
        dialog._btn_confirm_import.click()

    assert save_preview.called


def test_order_detail_import_button_opens_preview_dialog(session_factory) -> None:
    app()
    page = OrderDetailPage(order_service=OrderService(session_factory))

    assert page._btn_import_orders.isEnabled()
    assert page._btn_import_orders.text() == "导入订单"
    assert "预览" in page._btn_import_orders.toolTip()

    created = []

    class FakeDialog:
        def __init__(self, parent=None):
            created.append(parent)

        def exec(self):
            return 0

    with patch("ui.pages.order_detail_page.OrderImportDialog", FakeDialog):
        page._btn_import_orders.click()

    assert created == [page]
    assert "订单导入窗口已关闭" in page._status_label.text()

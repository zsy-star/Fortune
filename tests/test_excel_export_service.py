from __future__ import annotations

import re
from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from models import OperationLog, Order, SettlementRecord
from schemas.draw_schema import LotteryDrawCreate
from schemas.order_schema import OrderCreate, OrderItemCreate
from services.draw_service import DrawService
from services.excel_export_service import ExcelExportError, ExcelExportService
from services.log_service import LogService
from services.order_service import OrderService
from services.settlement_service import SettlementService


def create_order(
    service: OrderService,
    *,
    customer: str = "export-customer",
    region: str = "澳门",
    status: str | None = None,
    amount: str = "10",
    created_at: datetime | None = None,
):
    order = service.create_order(
        OrderCreate(
            customer_name=customer,
            channel="test-channel",
            region=region,
            raw_text=f"{customer} {amount}",
            source="test",
            items=[OrderItemCreate(bet_type="特码", selection="01", amount=amount)],
        )
    )
    if status is not None or created_at is not None:
        with service._session_factory() as session:
            saved = session.get(Order, order.id)
            assert saved is not None
            if status is not None:
                saved.status = status
            if created_at is not None:
                saved.created_at = created_at
                saved.updated_at = created_at
            session.commit()
    return order


def settle_order(session_factory, order_id: int, order_no: str, settled_at: datetime) -> None:
    detail = OrderService(session_factory).get_order(order_id)
    assert detail is not None
    draw = DrawService(session_factory).create_draw(
        LotteryDrawCreate(
            region=detail.region,
            issue_number=str(200000 + order_id),
            draw_date=settled_at.date(),
            regular_numbers=["02", "03", "04", "05", "06", "07"],
            special_number="01",
        )
    )
    SettlementService(session_factory).commit_order_settlement(order_id, draw.id)
    with session_factory() as session:
        order = session.get(Order, order_id)
        record = session.query(SettlementRecord).filter_by(order_id=order_id).one()
        saved_log = session.get(OperationLog, record.operation_log_id)
        assert order is not None
        assert saved_log is not None
        record.settled_at = settled_at
        order.updated_at = settled_at
        saved_log.created_at = settled_at
        session.commit()


def workbook_rows(path: Path) -> list[tuple]:
    workbook = load_workbook(path)
    sheet = workbook.active
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_empty_orders_export_creates_excel_with_headers(session_factory, tmp_path) -> None:
    result = ExcelExportService(session_factory, output_dir=tmp_path / "exports").export_orders()

    rows = workbook_rows(result.export_path)
    assert result.report_type == "orders"
    assert result.row_count == 0
    assert result.size_bytes > 0
    assert rows == [
        (
            "订单ID",
            "订单号",
            "客户",
            "渠道",
            "地区",
            "状态",
            "投注总额",
            "来源",
            "创建时间",
            "更新时间",
        )
    ]
    assert result.export_path.parent == tmp_path / "exports"


def test_orders_export_headers_row_count_and_timestamp(session_factory, tmp_path) -> None:
    service = OrderService(session_factory)
    order = create_order(service, customer="张三", amount="12.50")

    result = ExcelExportService(session_factory).export_orders(output_dir=tmp_path / "exports")

    rows = workbook_rows(result.export_path)
    assert result.row_count == 1
    assert re.match(r"orders_export_\d{8}_\d{6}\.xlsx", result.file_name)
    assert rows[1][0] == order.id
    assert rows[1][1] == order.order_no
    assert rows[1][2] == "张三"
    assert rows[1][5] == "active"
    assert rows[1][6] == 12.5


def test_orders_export_excludes_voided_by_default_and_can_include_voided(session_factory, tmp_path) -> None:
    service = OrderService(session_factory)
    active = create_order(service, customer="active", amount="10")
    voided = create_order(service, customer="voided", amount="20")
    service.void_order(voided.id, "export exclusion")
    export_service = ExcelExportService(session_factory, output_dir=tmp_path / "exports")

    default_result = export_service.export_orders()
    include_result = export_service.export_orders(include_voided=True)

    default_ids = {row[0] for row in workbook_rows(default_result.export_path)[1:]}
    include_rows = workbook_rows(include_result.export_path)[1:]
    include_ids = {row[0] for row in include_rows}
    assert default_ids == {active.id}
    assert include_ids == {active.id, voided.id}
    assert any(row[5] == "voided" for row in include_rows)


def test_orders_export_filters_by_date_region_status_and_keyword(session_factory, tmp_path) -> None:
    service = OrderService(session_factory)
    old_day = datetime.now() - timedelta(days=3)
    new_day = datetime.now()
    old_order = create_order(service, customer="old", region="澳门", amount="10", created_at=old_day)
    settled = create_order(
        service,
        customer="settled",
        region="香港",
        status="settled",
        amount="20",
        created_at=new_day,
    )
    create_order(service, customer="active-hk", region="香港", amount="30", created_at=new_day)
    export_service = ExcelExportService(session_factory, output_dir=tmp_path / "exports")

    filtered = export_service.export_orders(
        region="香港",
        status="settled",
        start_date=new_day - timedelta(minutes=1),
        end_date=new_day + timedelta(minutes=1),
    )
    keyword_by_no = export_service.export_orders(keyword=old_order.order_no)
    keyword_by_id = export_service.export_orders(keyword=str(settled.id))

    assert [row[0] for row in workbook_rows(filtered.export_path)[1:]] == [settled.id]
    assert [row[0] for row in workbook_rows(keyword_by_no.export_path)[1:]] == [old_order.id]
    assert [row[0] for row in workbook_rows(keyword_by_id.export_path)[1:]] == [settled.id]


def test_orders_export_filters_by_declarer(session_factory, tmp_path) -> None:
    service = OrderService(session_factory)
    target = create_order(service, customer="林林", region="澳门")
    create_order(service, customer="老汪", region="澳门")

    result = ExcelExportService(session_factory).export_orders(
        declarer_name="林林",
        output_dir=tmp_path,
    )

    assert [row[0] for row in workbook_rows(result.export_path)[1:]] == [target.id]
    assert result.filters["declarer_name"] == "林林"


def test_export_directory_is_created_and_real_exports_dir_not_used(session_factory, tmp_path) -> None:
    output_dir = tmp_path / "nested" / "exports"

    result = ExcelExportService(session_factory).export_orders(output_dir=output_dir)

    assert output_dir.exists()
    assert result.export_path.parent == output_dir
    assert "data/fortune.db" not in str(result.export_path).replace("\\", "/")


def test_settlement_ledger_exports_only_settled_and_not_voided(session_factory, tmp_path) -> None:
    service = OrderService(session_factory)
    settled = create_order(service, customer="settled", amount="10")
    active = create_order(service, customer="active", amount="20")
    voided = create_order(service, customer="voided", amount="30")
    service.void_order(voided.id, "not ledger")
    settle_order(session_factory, settled.id, settled.order_no, datetime.now())

    result = ExcelExportService(session_factory, output_dir=tmp_path / "exports").export_settlement_ledger()

    rows = workbook_rows(result.export_path)
    assert rows[0] == (
        "结算ID",
        "订单ID",
        "订单号",
        "客户",
        "地区",
        "状态",
        "投注总额",
        "结算时间",
        "开奖期号",
        "命中数量",
        "未中数量",
        "不支持数量",
        "最近操作日志摘要",
        "创建时间",
        "更新时间",
    )
    assert result.row_count == 1
    assert rows[1][1] == settled.id
    assert rows[1][5] == "settled"
    assert rows[1][8] != "-"
    assert rows[1][9] == 1
    assert rows[1][10] == 0
    assert rows[1][11] == 0
    assert settled.order_no in rows[1][12]
    assert active.id not in {row[1] for row in rows[1:]}
    assert voided.id not in {row[1] for row in rows[1:]}


def test_operation_logs_export_headers_row_count_and_filters(session_factory, tmp_path) -> None:
    log_service = LogService(session_factory)
    first = log_service.create_log(
        module="order",
        action="void",
        description="void success",
        operator="tester",
        related_type="order",
        related_id=101,
    )
    log_service.create_log(
        module="settlement",
        action="commit",
        description="settlement success",
        operator="system",
        related_type="order",
        related_id=202,
    )

    result = ExcelExportService(session_factory, output_dir=tmp_path / "exports").export_operation_logs(
        module="order",
        action="void",
        related_id=101,
        keyword="void",
    )

    rows = workbook_rows(result.export_path)
    assert rows[0] == (
        "日志ID",
        "时间",
        "模块",
        "动作",
        "操作人",
        "对象类型",
        "对象ID",
        "结果",
        "摘要 / 详情",
    )
    assert result.row_count == 1
    assert rows[1][0] == first.id
    assert rows[1][2] == "order"
    assert rows[1][3] == "void"
    assert rows[1][4] == "tester"
    assert rows[1][6] == 101
    assert rows[1][7] == "成功"


def test_exports_are_read_only_for_database(session_factory, tmp_path) -> None:
    order_service = OrderService(session_factory)
    log_service = LogService(session_factory)
    create_order(order_service)
    before_orders = order_service.count_orders()
    before_logs = log_service.count_logs()

    export_service = ExcelExportService(session_factory, output_dir=tmp_path / "exports")
    export_service.export_orders()
    export_service.export_operation_logs()

    assert order_service.count_orders() == before_orders
    assert log_service.count_logs() == before_logs


def test_write_failure_removes_partial_temp_file(monkeypatch, session_factory, tmp_path) -> None:
    output_dir = tmp_path / "exports"

    def fail_save(self, filename):
        Path(filename).write_bytes(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr("services.excel_export_service.Workbook.save", fail_save)

    with pytest.raises(ExcelExportError, match="disk full"):
        ExcelExportService(session_factory, output_dir=output_dir).export_orders()

    assert list(output_dir.glob("*.xlsx")) == []
    assert list(output_dir.glob("*.tmp")) == []

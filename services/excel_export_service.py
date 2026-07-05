"""Read-only Excel export service."""

from __future__ import annotations

from collections.abc import Callable, Iterable
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from core.database import SessionLocal
from schemas.excel_export_schema import ExcelExportResult
from schemas.log_schema import OperationLogResult
from schemas.order_schema import OrderSummary
from schemas.settlement_schema import SettlementLedgerResult
from services.log_service import LogService
from services.order_service import OrderService

DEFAULT_EXPORT_DIR = Path("exports")
ORDER_STATUS_VOIDED = "voided"


class ExcelExportError(RuntimeError):
    """Raised when an export cannot be completed safely."""


class ExcelExportService:
    def __init__(
        self,
        session_factory: Callable[[], Session] = SessionLocal,
        *,
        output_dir: str | Path | None = None,
    ):
        self._session_factory = session_factory
        self._order_service = OrderService(session_factory)
        self._log_service = LogService(session_factory)
        self._output_dir = Path(output_dir) if output_dir is not None else DEFAULT_EXPORT_DIR

    def export_orders(
        self,
        *,
        region: str | None = None,
        status: str | None = None,
        start_date: datetime | None = None,
        end_date: datetime | None = None,
        keyword: str | None = None,
        declarer_name: str | None = None,
        bet_type: str | None = None,
        winning_status: str | None = None,
        include_voided: bool = False,
        output_dir: str | Path | None = None,
    ) -> ExcelExportResult:
        filters = {
            "region": region,
            "status": status,
            "start_date": start_date,
            "end_date": end_date,
            "keyword": keyword,
            "declarer_name": declarer_name,
            "bet_type": bet_type,
            "winning_status": winning_status,
            "include_voided": include_voided,
        }
        orders = self._list_orders_for_export(
            region=region,
            status=status,
            start_date=start_date,
            end_date=end_date,
            keyword=keyword,
            declarer_name=declarer_name,
            bet_type=bet_type,
            winning_status=winning_status,
            include_voided=include_voided,
        )
        headers = [
            "订单ID",
            "订单号",
            "客户",
            "渠道",
            "地区",
            "状态",
            "投注总额",
            "来源",
            "生肖年份",
            "创建时间",
            "更新时间",
        ]
        rows = [
            [
                order.id,
                order.order_no,
                _dash(order.customer_name),
                _dash(order.channel),
                order.region,
                order.status,
                _decimal_to_float(order.total_amount),
                _dash(order.source),
                order.zodiac_year if order.zodiac_year is not None else "-",
                order.created_at,
                order.updated_at,
            ]
            for order in orders
        ]
        return self._export_workbook(
            report_type="orders",
            file_prefix="orders_export",
            sheet_name="订单列表",
            headers=headers,
            rows=rows,
            filters=filters,
            output_dir=output_dir,
        )

    def export_settlement_ledger(
        self,
        *,
        region: str | None = None,
        keyword: str | None = None,
        start_date: datetime | None = None,
        end_date: datetime | None = None,
        output_dir: str | Path | None = None,
    ) -> ExcelExportResult:
        filters = {
            "region": region,
            "keyword": keyword,
            "start_date": start_date,
            "end_date": end_date,
        }
        orders = self._list_settlement_ledger_for_export(
            region=region,
            keyword=keyword,
            start_date=start_date,
            end_date=end_date,
        )
        headers = [
            "结算ID",
            "订单ID",
            "订单号",
            "客户",
            "地区",
            "状态",
            "投注总额",
            "结算时间",
            "开奖期号",
            "命中数量",
            "未中数量",
            "不支持数量",
            "最近操作日志摘要",
            "创建时间",
            "更新时间",
        ]
        rows = []
        for record in orders:
            rows.append(
                [
                    record.id,
                    record.order_id,
                    record.order_no,
                    _dash(record.customer_name),
                    record.region,
                    record.order_status,
                    _decimal_to_float(record.total_amount),
                    record.settled_at,
                    record.issue_number,
                    record.hit_count,
                    record.miss_count,
                    record.unsupported_count,
                    record.operation_log_description or "-",
                    record.order_created_at,
                    record.order_updated_at,
                ]
            )
        return self._export_workbook(
            report_type="settlement_ledger",
            file_prefix="settlement_ledger_export",
            sheet_name="结算流水",
            headers=headers,
            rows=rows,
            filters=filters,
            output_dir=output_dir,
        )

    def export_operation_logs(
        self,
        *,
        module: str | None = None,
        action: str | None = None,
        operator: str | None = None,
        related_type: str | None = None,
        related_id: int | None = None,
        start_date: datetime | None = None,
        end_date: datetime | None = None,
        keyword: str | None = None,
        output_dir: str | Path | None = None,
    ) -> ExcelExportResult:
        filters = {
            "module": module,
            "action": action,
            "operator": operator,
            "related_type": related_type,
            "related_id": related_id,
            "start_date": start_date,
            "end_date": end_date,
            "keyword": keyword,
        }
        logs = self._list_logs_for_export(
            module=module,
            action=action,
            operator=operator,
            related_type=related_type,
            related_id=related_id,
            start_date=start_date,
            end_date=end_date,
            keyword=keyword,
        )
        headers = [
            "日志ID",
            "时间",
            "模块",
            "动作",
            "操作人",
            "对象类型",
            "对象ID",
            "结果",
            "摘要 / 详情",
        ]
        rows = [
            [
                log.id,
                log.created_at,
                log.module,
                log.action,
                _dash(log.operator),
                _dash(log.related_type),
                log.related_id if log.related_id is not None else "-",
                _result_from_description(log.description),
                log.description,
            ]
            for log in logs
        ]
        return self._export_workbook(
            report_type="operation_logs",
            file_prefix="operation_logs_export",
            sheet_name="操作日志",
            headers=headers,
            rows=rows,
            filters=filters,
            output_dir=output_dir,
        )

    def _list_orders_for_export(
        self,
        *,
        region: str | None,
        status: str | None,
        start_date: datetime | None,
        end_date: datetime | None,
        keyword: str | None,
        declarer_name: str | None,
        bet_type: str | None,
        winning_status: str | None,
        include_voided: bool,
    ) -> list[OrderSummary]:
        keyword = keyword.strip() if keyword else None
        if not include_voided and status == ORDER_STATUS_VOIDED:
            return []
        if keyword and keyword.isdigit():
            detail = self._order_service.get_order(int(keyword))
            if detail is None:
                return []
            summary = OrderSummary(
                id=detail.id,
                order_no=detail.order_no,
                customer_name=detail.customer_name,
                channel=detail.channel,
                region=detail.region,
                source=detail.source,
                raw_text=detail.raw_text,
                total_amount=detail.total_amount,
                status=detail.status,
                created_at=detail.created_at,
                updated_at=detail.updated_at,
                item_count=len(detail.items),
                zodiac_year=detail.zodiac_year,
            )
            return [
                summary
            ] if self._matches_order_filters(
                summary,
                region=region,
                status=status,
                start_date=start_date,
                end_date=end_date,
                declarer_name=declarer_name,
                bet_type=bet_type,
                winning_status=winning_status,
                include_voided=include_voided,
            ) else []

        rows = self._page_orders(
            region=region,
            order_no=keyword,
            declarer_name=declarer_name,
            bet_type=bet_type,
            winning_status=winning_status,
            status=status,
            start_date=start_date,
            end_date=end_date,
        )
        if include_voided:
            return rows
        return [order for order in rows if order.status != ORDER_STATUS_VOIDED]

    def _list_settlement_ledger_for_export(
        self,
        *,
        region: str | None,
        keyword: str | None,
        start_date: datetime | None,
        end_date: datetime | None,
    ) -> list[SettlementLedgerResult]:
        rows: list[SettlementLedgerResult] = []
        offset = 0
        page_size = 200
        while True:
            page = self._order_service.list_settlement_ledger(
                region=region,
                keyword=keyword,
                start_date=start_date,
                end_date=end_date,
                limit=page_size,
                offset=offset,
            )
            rows.extend(page)
            if len(page) < page_size:
                break
            offset += page_size
        return rows

    def _list_logs_for_export(
        self,
        *,
        module: str | None,
        action: str | None,
        operator: str | None,
        related_type: str | None,
        related_id: int | None,
        start_date: datetime | None,
        end_date: datetime | None,
        keyword: str | None,
    ) -> list[OperationLogResult]:
        rows: list[OperationLogResult] = []
        offset = 0
        page_size = 500
        while True:
            page = self._log_service.list_logs(
                module=module,
                action=action,
                operator=operator,
                related_type=related_type,
                related_id=related_id,
                start_date=start_date,
                end_date=end_date,
                keyword=keyword,
                limit=page_size,
                offset=offset,
            )
            rows.extend(page)
            if len(page) < page_size:
                break
            offset += page_size
        return rows

    def _page_orders(
        self,
        *,
        region: str | None,
        order_no: str | None,
        declarer_name: str | None,
        bet_type: str | None,
        winning_status: str | None,
        status: str | None,
        start_date: datetime | None,
        end_date: datetime | None,
    ) -> list[OrderSummary]:
        rows: list[OrderSummary] = []
        offset = 0
        page_size = 200
        while True:
            page = self._order_service.list_orders(
                region=region,
                order_no=order_no,
                declarer_name=declarer_name,
                bet_type=bet_type,
                winning_status=winning_status,
                status=status,
                start_date=start_date,
                end_date=end_date,
                limit=page_size,
                offset=offset,
            )
            rows.extend(page)
            if len(page) < page_size:
                break
            offset += page_size
        return rows

    def _latest_settlement_log(self, order: OrderSummary) -> OperationLogResult | None:
        logs = self._log_service.list_logs(
            module="settlement",
            action="commit",
            related_type="order",
            related_id=order.id,
            limit=1,
        )
        if logs:
            return logs[0]
        fallback = self._log_service.list_logs(
            module="settlement",
            action="commit",
            related_type="order",
            keyword=order.order_no,
            limit=1,
        )
        return fallback[0] if fallback else None

    def _matches_order_filters(
        self,
        order: OrderSummary,
        *,
        region: str | None,
        status: str | None,
        start_date: datetime | None,
        end_date: datetime | None,
        declarer_name: str | None,
        bet_type: str | None,
        winning_status: str | None,
        include_voided: bool,
    ) -> bool:
        if region and order.region != region:
            return False
        if status and order.status != status:
            return False
        if declarer_name and order.customer_name != declarer_name:
            return False
        if bet_type:
            detail = self._order_service.get_order(order.id)
            if detail is None or all(item.bet_type != bet_type for item in detail.items):
                return False
        if winning_status:
            matched = self._order_service.count_orders(
                order_no=order.order_no,
                winning_status=winning_status,
            )
            if matched <= 0:
                return False
        if not include_voided and order.status == ORDER_STATUS_VOIDED:
            return False
        if start_date and order.created_at < start_date:
            return False
        if end_date and order.created_at > end_date:
            return False
        return True

    def _export_workbook(
        self,
        *,
        report_type: str,
        file_prefix: str,
        sheet_name: str,
        headers: list[str],
        rows: list[list[Any]],
        filters: dict[str, Any],
        output_dir: str | Path | None,
    ) -> ExcelExportResult:
        created_at = datetime.now()
        target_dir = Path(output_dir) if output_dir is not None else self._output_dir
        target_dir.mkdir(parents=True, exist_ok=True)
        export_path = self._unique_export_path(target_dir, file_prefix, created_at)
        temp_path = export_path.with_name(f".{export_path.name}.tmp")

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        self._format_sheet(worksheet, headers, rows)

        try:
            workbook.save(temp_path)
            temp_path.replace(export_path)
        except Exception as exc:
            if temp_path.exists():
                temp_path.unlink()
            if export_path.exists() and export_path.stat().st_size == 0:
                export_path.unlink()
            raise ExcelExportError(f"Excel export failed: {exc}") from exc

        size_bytes = export_path.stat().st_size
        if size_bytes <= 0:
            raise ExcelExportError(f"Excel export file is empty: {export_path}")
        return ExcelExportResult(
            export_path=export_path,
            file_name=export_path.name,
            report_type=report_type,
            row_count=len(rows),
            created_at=created_at,
            size_bytes=size_bytes,
            filters=_serializable_filters(filters),
        )

    def _unique_export_path(self, target_dir: Path, file_prefix: str, created_at: datetime) -> Path:
        timestamp = created_at.strftime("%Y%m%d_%H%M%S")
        base_name = f"{file_prefix}_{timestamp}"
        candidate = target_dir / f"{base_name}.xlsx"
        suffix = 1
        while candidate.exists():
            candidate = target_dir / f"{base_name}_{suffix:03d}.xlsx"
            suffix += 1
        return candidate

    def _format_sheet(self, worksheet, headers: list[str], rows: list[list[Any]]) -> None:
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.00"

        for idx, header in enumerate(headers, start=1):
            values: Iterable[Any] = [header]
            if rows:
                values = [header, *(row[idx - 1] for row in rows)]
            width = min(max(max(len(str(value)) for value in values if value is not None) + 2, 10), 48)
            worksheet.column_dimensions[get_column_letter(idx)].width = width


def _dash(value: object | None) -> str:
    return str(value) if value not in (None, "") else "-"


def _decimal_to_float(value: Decimal) -> float:
    return float(value)


def _result_from_description(description: str) -> str:
    return "成功" if "success" in description.lower() else "-"


def _serializable_filters(filters: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in filters.items():
        result[key] = value.isoformat(sep=" ") if isinstance(value, datetime) else value
    return result
